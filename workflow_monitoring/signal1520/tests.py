import json

from django.contrib.auth.models import User, Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from authentication.models import Profile
from .models import (
    AlarmInfo, Attachment, Comment, Knowledge, Organization,
    Road, Station, System, Task, UserKnowledge,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def make_user(username, password='testpass', org=None, codenames=()):
    """Создаёт пользователя с Profile и опциональным набором прав."""
    user = User.objects.create_user(username, password=password)
    for codename in codenames:
        user.user_permissions.add(Permission.objects.get(codename=codename))
    Profile.objects.create(user=user, organization=org)
    return user


# ---------------------------------------------------------------------------
# Model tests
# ---------------------------------------------------------------------------

class RoadModelTest(TestCase):
    def test_str(self):
        """__str__ возвращает название дороги — используется в формах и шаблонах."""
        road = Road(title='Московская дорога')
        self.assertEqual(str(road), 'Московская дорога')

    def test_ordering(self):
        """Дороги упорядочены по алфавиту по полю title (Meta.ordering=['title'])."""
        Road.objects.create(title='Ярославская')
        Road.objects.create(title='Арктическая')
        titles = list(Road.objects.values_list('title', flat=True))
        self.assertEqual(titles, sorted(titles))


class SystemModelTest(TestCase):
    def test_str(self):
        """__str__ возвращает название системы — используется в формах и шаблонах."""
        system = System(title='АРС-ДП')
        self.assertEqual(str(system), 'АРС-ДП')

    def test_ordering(self):
        """Системы упорядочены по алфавиту по полю title (Meta.ordering=['title'])."""
        System.objects.create(title='Система Б')
        System.objects.create(title='Система А')
        titles = list(System.objects.values_list('title', flat=True))
        self.assertEqual(titles, sorted(titles))


class StationModelTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('u', password='p')
        cls.road = Road.objects.create(title='Московская дорога')
        cls.system = System.objects.create(title='АРС-ДП')

    def test_str_with_road(self):
        """__str__ включает имя станции и название дороги через FK Road.__str__."""
        station = Station(name='Тест', road=self.road)
        self.assertIn('Тест', str(station))
        self.assertIn('Московская дорога', str(station))

    def test_system_nullable(self):
        """Поле system на станции nullable — станцию можно создать без указания системы."""
        station = Station.objects.create(
            name='Без системы', road=self.road, created_by=self.user
        )
        self.assertIsNone(station.system)

    def test_ordering_by_road_title(self):
        """Станции сортируются по названию дороги (road__title), а не по FK-id."""
        road2 = Road.objects.create(title='Арктическая дорога')
        Station.objects.create(name='С1', road=self.road, created_by=self.user)
        Station.objects.create(name='С2', road=road2, created_by=self.user)
        first = Station.objects.first()
        self.assertEqual(first.road.title, 'Арктическая дорога')


class TaskModelTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('u', password='p')
        cls.road = Road.objects.create(title='Дорога')
        cls.station = Station.objects.create(
            name='Станция', road=cls.road, created_by=cls.user
        )

    def test_default_status_is_new(self):
        """Новая задача получает статус 'new' по умолчанию без явного указания."""
        task = Task.objects.create(station=self.station, description='Задача')
        self.assertEqual(task.status, Task.Status.NEW)

    def test_str(self):
        """__str__ содержит id задачи и имя станции для удобной идентификации в admin."""
        task = Task.objects.create(station=self.station, description='Задача')
        self.assertIn(str(task.id), str(task))
        self.assertIn(self.station.name, str(task))

    def test_all_statuses_valid(self):
        """Task.Status содержит все четыре ожидаемых значения: new, in_progress, completed, cancelled."""
        valid = {s[0] for s in Task.Status.choices}
        self.assertIn('new', valid)
        self.assertIn('in_progress', valid)
        self.assertIn('completed', valid)
        self.assertIn('cancelled', valid)


class CommentModelTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('u', first_name='Иван', last_name='Петров', password='p')
        cls.road = Road.objects.create(title='Дорога')
        cls.station = Station.objects.create(name='Ст', road=cls.road, created_by=cls.user)
        cls.task = Task.objects.create(station=cls.station, description='Задача')

    def test_author_name_full_name(self):
        """author_name() возвращает 'Имя Фамилия', если оба поля заполнены у пользователя."""
        comment = Comment(task=self.task, user=self.user, body='Текст')
        self.assertEqual(comment.author_name(), 'Иван Петров')

    def test_author_name_username_fallback(self):
        """author_name() возвращает username, если first_name и last_name не заполнены."""
        user2 = User.objects.create_user('only_username', password='p')
        comment = Comment(task=self.task, user=user2, body='Текст')
        self.assertEqual(comment.author_name(), 'only_username')

    def test_author_name_anon(self):
        """author_name() возвращает 'Аноним', если user=None (удалённый или незалогиненный)."""
        comment = Comment(task=self.task, user=None, body='Текст')
        self.assertEqual(comment.author_name(), 'Аноним')


class AttachmentModelTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('u', password='p')
        cls.road = Road.objects.create(title='Дорога')
        cls.station = Station.objects.create(name='Ст', road=cls.road, created_by=cls.user)
        cls.task = Task.objects.create(station=cls.station, description='Задача')

    def _make_attachment(self, filename):
        f = SimpleUploadedFile(filename, b'data')
        return Attachment(task=self.task, file=f)

    def test_is_image_jpg(self):
        """is_image() возвращает True для файлов с расширением .jpg."""
        a = self._make_attachment('photo.jpg')
        self.assertTrue(a.is_image())

    def test_is_image_png(self):
        """is_image() возвращает True для файлов с расширением .png."""
        a = self._make_attachment('photo.png')
        self.assertTrue(a.is_image())

    def test_is_image_false_for_pdf(self):
        """is_image() возвращает False для не-графических файлов, например .pdf."""
        a = self._make_attachment('document.pdf')
        self.assertFalse(a.is_image())

    def test_filename(self):
        """filename() возвращает чистое имя файла без пути — используется в шаблоне вложений."""
        a = self._make_attachment('myfile.txt')
        self.assertEqual(a.filename(), 'myfile.txt')


# ---------------------------------------------------------------------------
# View tests — base setup
# ---------------------------------------------------------------------------

ORG_SLUG = 'testorg'


class ViewTestBase(TestCase):
    """Создаёт организацию, пользователей с профилями и фикстуры для тестов views."""

    @classmethod
    def setUpTestData(cls):
        cls.org = Organization.objects.create(name='Test Org', slug=ORG_SLUG)

        cls.superuser = User.objects.create_superuser('admin', password='pass')
        Profile.objects.create(user=cls.superuser, organization=cls.org)

        cls.plain_user = User.objects.create_user('plain', password='pass')
        Profile.objects.create(user=cls.plain_user, organization=cls.org)

        station_perms = ['view_station', 'add_station', 'change_station']
        task_perms = ['view_task', 'add_task', 'change_task']
        cls.station_user = make_user('station_u', org=cls.org, codenames=station_perms)
        cls.task_user = make_user('task_u', org=cls.org, codenames=task_perms + station_perms)

        cls.road = Road.objects.create(title='Тестовая дорога', organization=cls.org)
        cls.system = System.objects.create(title='Тестовая система', organization=cls.org)
        cls.station = Station.objects.create(
            name='Тест Станция',
            road=cls.road,
            system=cls.system,
            distance='ДЦС-1',
            description='Описание тестовой станции',
            created_by=cls.superuser,
            organization=cls.org,
        )
        cls.task = Task.objects.create(
            station=cls.station,
            description='Тестовая задача',
            status=Task.Status.NEW,
            responsible_user=cls.plain_user,
        )
        cls.alarm = AlarmInfo.objects.create(
            number='T001',
            description='Тестовый аларм описание',
            explanation='Пояснение к аларму',
        )

    def url(self, name, **kwargs):
        """Строит URL с org_slug='testorg' плюс дополнительные kwargs."""
        return reverse(name, kwargs={'org_slug': ORG_SLUG, **kwargs})


# ---------------------------------------------------------------------------
# OrgMixin — изоляция по организации
# ---------------------------------------------------------------------------

class OrgIsolationTest(ViewTestBase):
    def test_user_from_other_org_gets_403(self):
        """Пользователь из другой организации не может открыть страницу этой орги — 403."""
        other_org = Organization.objects.create(name='Other', slug='other')
        other_user = make_user('other_u', org=other_org, codenames=['view_station'])
        self.client.force_login(other_user)
        r = self.client.get(self.url('signal1520:station_list'))
        self.assertEqual(r.status_code, 403)

    def test_anon_gets_login_redirect(self):
        """Анонимный пользователь перенаправляется на логин, а не получает 403."""
        r = self.client.get(self.url('signal1520:station_list'))
        self.assertEqual(r.status_code, 302)
        self.assertIn('/accounts/login/', r['Location'])


# ---------------------------------------------------------------------------
# TasksIndexView
# ---------------------------------------------------------------------------

class TasksIndexViewTest(ViewTestBase):
    def test_redirect_if_not_logged_in(self):
        """Анонимный GET на главную перенаправляет на страницу логина (LoginRequiredMixin)."""
        r = self.client.get(self.url('signal1520:index'))
        self.assertEqual(r.status_code, 302)
        self.assertIn('/accounts/login/', r['Location'])

    def test_logged_in_returns_200(self):
        """Любой авторизованный пользователь своей орги видит главную страницу."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:index'))
        self.assertEqual(r.status_code, 200)


# ---------------------------------------------------------------------------
# StationListView
# ---------------------------------------------------------------------------

class StationListViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET перенаправляет на логин с параметром next."""
        u = self.url('signal1520:station_list')
        r = self.client.get(u)
        self.assertRedirects(r, f'/accounts/login/?next={u}')

    def test_no_permission_returns_403(self):
        """Авторизованный пользователь без view_station получает 403."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:station_list'))
        self.assertEqual(r.status_code, 403)

    def test_with_permission_returns_200(self):
        """Пользователь с правом view_station видит список станций."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_list'))
        self.assertEqual(r.status_code, 200)

    def test_superuser_returns_200(self):
        """Суперпользователь всегда проходит test_func и видит список станций."""
        self.client.force_login(self.superuser)
        r = self.client.get(self.url('signal1520:station_list'))
        self.assertEqual(r.status_code, 200)

    def test_station_appears_in_list(self):
        """Существующая станция отображается в таблице на странице списка."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_list'))
        self.assertContains(r, 'Тест Станция')

    def test_search_by_name(self):
        """Поиск по точному имени станции возвращает эту станцию в результатах."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_list'), {'search': 'Тест Станция'})
        self.assertContains(r, 'Тест Станция')

    def test_search_by_road(self):
        """Поиск по названию дороги (road__title) находит станции этой дороги."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_list'), {'search': 'Тестовая дорога'})
        self.assertContains(r, 'Тест Станция')

    def test_search_by_system(self):
        """Поиск по названию системы (system__title) находит станции с этой системой."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_list'), {'search': 'Тестовая система'})
        self.assertContains(r, 'Тест Станция')

    def test_search_no_results(self):
        """Поиск по несуществующей строке не возвращает ни одной станции."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_list'), {'search': 'несуществующий_xyz'})
        self.assertNotContains(r, 'Тест Станция')


# ---------------------------------------------------------------------------
# StationDetailView
# ---------------------------------------------------------------------------

class StationDetailViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на страницу станции перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:station_details', pk=self.station.pk))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_returns_403(self):
        """Авторизованный пользователь без view_station получает 403."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:station_details', pk=self.station.pk))
        self.assertEqual(r.status_code, 403)

    def test_with_permission_returns_200(self):
        """Пользователь с view_station открывает страницу детали станции."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_details', pk=self.station.pk))
        self.assertEqual(r.status_code, 200)

    def test_shows_station_data(self):
        """Страница детали содержит имя, дорогу и систему станции."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_details', pk=self.station.pk))
        self.assertContains(r, 'Тест Станция')
        self.assertContains(r, 'Тестовая дорога')
        self.assertContains(r, 'Тестовая система')

    def test_shows_task_in_detail(self):
        """Страница детали станции показывает привязанные к ней задачи."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_details', pk=self.station.pk))
        self.assertContains(r, 'Тестовая задача')

    def test_404_for_nonexistent_station(self):
        """Запрос на несуществующий pk станции возвращает 404."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_details', pk=99999))
        self.assertEqual(r.status_code, 404)


# ---------------------------------------------------------------------------
# StationCreateView
# ---------------------------------------------------------------------------

class StationCreateViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на форму создания станции перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:create_station'))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_redirects_to_error(self):
        """Пользователь без add_station перенаправляется на страницу ошибки."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:create_station'))
        self.assertEqual(r.status_code, 302)
        self.assertIn('error', r['Location'])

    def test_form_renders_for_permitted_user(self):
        """Пользователь с add_station видит форму создания станции (GET 200)."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:create_station'))
        self.assertEqual(r.status_code, 200)

    def test_form_contains_only_org_roads(self):
        """Выпадающий список road содержит только дороги своей организации."""
        other_org = Organization.objects.create(name='Other', slug='other2')
        Road.objects.create(title='Чужая дорога', organization=other_org)
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:create_station'))
        self.assertContains(r, 'Тестовая дорога')
        self.assertNotContains(r, 'Чужая дорога')

    def test_create_station_post(self):
        """POST с валидными данными создаёт новую станцию и редиректит на её детальную страницу."""
        self.client.force_login(self.station_user)
        r = self.client.post(self.url('signal1520:create_station'), {
            'name': 'Новая станция',
            'road': self.road.pk,
            'system': self.system.pk,
            'distance': 'ДЦС-2',
            'description': 'Описание',
            'latitude': '',
            'longitude': '',
        })
        self.assertEqual(r.status_code, 302)
        self.assertTrue(Station.objects.filter(name='Новая станция').exists())

    def test_create_station_sets_created_by(self):
        """form_valid() автоматически записывает в created_by текущего пользователя."""
        self.client.force_login(self.station_user)
        self.client.post(self.url('signal1520:create_station'), {
            'name': 'Станция пользователя',
            'road': self.road.pk,
            'distance': '',
            'description': '',
            'latitude': '',
            'longitude': '',
        })
        station = Station.objects.get(name='Станция пользователя')
        self.assertEqual(station.created_by, self.station_user)

    def test_create_station_sets_org(self):
        """form_valid() автоматически записывает organization из URL."""
        self.client.force_login(self.station_user)
        self.client.post(self.url('signal1520:create_station'), {
            'name': 'Станция с оргой',
            'road': self.road.pk,
            'distance': '',
            'description': '',
            'latitude': '',
            'longitude': '',
        })
        station = Station.objects.get(name='Станция с оргой')
        self.assertEqual(station.organization, self.org)

    def test_create_station_without_required_fields_fails(self):
        """POST без обязательных полей возвращает форму с ошибками, объект не создаётся."""
        self.client.force_login(self.station_user)
        r = self.client.post(self.url('signal1520:create_station'), {'description': 'Без названия'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(Station.objects.filter(description='Без названия').exists())


# ---------------------------------------------------------------------------
# StationUpdateView
# ---------------------------------------------------------------------------

class StationUpdateViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на форму редактирования станции перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:station_update', pk=self.station.pk))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_returns_403(self):
        """Авторизованный пользователь без change_station получает 403."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:station_update', pk=self.station.pk))
        self.assertEqual(r.status_code, 403)

    def test_form_renders_for_permitted_user(self):
        """Пользователь с change_station видит форму редактирования (GET 200)."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:station_update', pk=self.station.pk))
        self.assertEqual(r.status_code, 200)

    def test_update_station(self):
        """POST с новым именем обновляет станцию в БД."""
        self.client.force_login(self.station_user)
        r = self.client.post(self.url('signal1520:station_update', pk=self.station.pk), {
            'name': 'Обновлённое название',
            'road': self.road.pk,
            'system': self.system.pk,
            'distance': 'ДЦС-9',
            'description': '',
            'latitude': '',
            'longitude': '',
        })
        self.assertEqual(r.status_code, 302)
        self.station.refresh_from_db()
        self.assertEqual(self.station.name, 'Обновлённое название')

    def test_superuser_can_update(self):
        """Суперпользователь проходит test_func и получает форму редактирования."""
        self.client.force_login(self.superuser)
        r = self.client.get(self.url('signal1520:station_update', pk=self.station.pk))
        self.assertEqual(r.status_code, 200)


# ---------------------------------------------------------------------------
# BugsListView
# ---------------------------------------------------------------------------

class BugsListViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на список задач перенаправляет на логин с параметром next."""
        u = self.url('signal1520:bugs_list')
        r = self.client.get(u)
        self.assertRedirects(r, f'/accounts/login/?next={u}')

    def test_no_permission_returns_403(self):
        """Авторизованный пользователь без view_task получает 403."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:bugs_list'))
        self.assertEqual(r.status_code, 403)

    def test_with_permission_returns_200(self):
        """Пользователь с view_task видит список задач."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bugs_list'))
        self.assertEqual(r.status_code, 200)

    def test_task_appears_in_list(self):
        """Существующая задача отображается в таблице списка."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bugs_list'))
        self.assertContains(r, 'Тестовая задача')

    def test_search_by_description(self):
        """Поиск по тексту описания задачи находит нужную запись."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bugs_list'), {'search': 'Тестовая задача'})
        self.assertContains(r, 'Тестовая задача')

    def test_search_by_station_name(self):
        """Поиск по имени станции (station__name) находит задачи этой станции."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bugs_list'), {'search': 'Тест Станция'})
        self.assertContains(r, 'Тестовая задача')

    def test_search_no_result(self):
        """Поиск по несуществующей строке не возвращает задач."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bugs_list'), {'search': 'нет_такого_xyz'})
        self.assertNotContains(r, 'Тестовая задача')


# ---------------------------------------------------------------------------
# MyBugsListView
# ---------------------------------------------------------------------------

class MyBugsListViewTest(ViewTestBase):
    def setUp(self):
        self.responsible_user = make_user(
            'responsible', org=self.org, codenames=['view_task']
        )
        self.own_task = Task.objects.create(
            station=self.station,
            description='Задача resp_user',
            responsible_user=self.responsible_user,
        )

    def test_redirect_anon(self):
        """Анонимный GET на «мои задачи» перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:bugs_list_my'))
        self.assertEqual(r.status_code, 302)

    def test_shows_only_own_tasks(self):
        """Пользователь видит только задачи, где он назначен ответственным."""
        self.client.force_login(self.responsible_user)
        r = self.client.get(self.url('signal1520:bugs_list_my'))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Задача resp_user')

    def test_does_not_show_others_tasks(self):
        """Задачи, назначенные другому пользователю, не попадают в список 'мои'."""
        self.client.force_login(self.responsible_user)
        r = self.client.get(self.url('signal1520:bugs_list_my'))
        self.assertNotContains(r, 'Тестовая задача')

    def test_task_user_sees_only_own(self):
        """Пользователь без назначенных задач видит пустой список."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bugs_list_my'))
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, 'Задача resp_user')


# ---------------------------------------------------------------------------
# BugDetailView — GET
# ---------------------------------------------------------------------------

class BugDetailViewGetTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на страницу задачи перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:bug_details', pk=self.task.pk))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_returns_403(self):
        """Авторизованный пользователь без view_task получает 403."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:bug_details', pk=self.task.pk))
        self.assertEqual(r.status_code, 403)

    def test_with_permission_returns_200(self):
        """Пользователь с view_task открывает детальную страницу задачи."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bug_details', pk=self.task.pk))
        self.assertEqual(r.status_code, 200)

    def test_shows_task_description(self):
        """Страница содержит описание задачи."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bug_details', pk=self.task.pk))
        self.assertContains(r, 'Тестовая задача')

    def test_404_nonexistent_task(self):
        """Запрос на несуществующий pk задачи возвращает 404."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bug_details', pk=99999))
        self.assertEqual(r.status_code, 404)


# ---------------------------------------------------------------------------
# BugDetailView — POST: comment
# ---------------------------------------------------------------------------

class BugDetailPostCommentTest(ViewTestBase):
    def test_add_comment(self):
        """POST с comment_text создаёт новый Comment и редиректит обратно на задачу."""
        self.client.force_login(self.task_user)
        r = self.client.post(
            self.url('signal1520:bug_details', pk=self.task.pk),
            {'comment_text': 'Мой комментарий'},
        )
        self.assertEqual(r.status_code, 302)
        self.assertTrue(Comment.objects.filter(body='Мой комментарий', task=self.task).exists())

    def test_empty_comment_not_saved(self):
        """POST с пустой строкой не создаёт Comment в БД."""
        self.client.force_login(self.task_user)
        before = Comment.objects.count()
        self.client.post(
            self.url('signal1520:bug_details', pk=self.task.pk),
            {'comment_text': '   '},
        )
        self.assertEqual(Comment.objects.count(), before)

    def test_comment_sets_correct_user(self):
        """Созданный комментарий привязывается к текущему авторизованному пользователю."""
        self.client.force_login(self.task_user)
        self.client.post(
            self.url('signal1520:bug_details', pk=self.task.pk),
            {'comment_text': 'Проверка автора'},
        )
        comment = Comment.objects.get(body='Проверка автора')
        self.assertEqual(comment.user, self.task_user)


# ---------------------------------------------------------------------------
# BugDetailView — POST: status change (JSON)
# ---------------------------------------------------------------------------

class BugDetailPostStatusTest(ViewTestBase):
    def _patch_status(self, user, status):
        self.client.force_login(user)
        return self.client.post(
            self.url('signal1520:bug_details', pk=self.task.pk),
            data=json.dumps({'status': status}),
            content_type='application/json',
        )

    def test_superuser_can_change_status(self):
        """Суперпользователь меняет статус задачи через JSON POST."""
        r = self._patch_status(self.superuser, 'in_progress')
        self.assertEqual(r.status_code, 200)
        self.assertTrue(json.loads(r.content)['success'])
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, 'in_progress')

    def test_responsible_user_can_change_status(self):
        """Ответственный исполнитель с view_task может менять статус своей задачи."""
        responsible = make_user('resp_view', org=self.org, codenames=['view_task'])
        own_task = Task.objects.create(
            station=self.station,
            description='Задача для resp_view',
            responsible_user=responsible,
        )
        self.client.force_login(responsible)
        r = self.client.post(
            self.url('signal1520:bug_details', pk=own_task.pk),
            data=json.dumps({'status': 'completed'}),
            content_type='application/json',
        )
        self.assertEqual(r.status_code, 200)
        own_task.refresh_from_db()
        self.assertEqual(own_task.status, 'completed')

    def test_user_with_change_task_perm_can_change_status(self):
        """Пользователь с правом change_task может менять статус любой задачи."""
        r = self._patch_status(self.task_user, 'cancelled')
        self.assertEqual(r.status_code, 200)

    def test_invalid_status_returns_400(self):
        """JSON POST с недопустимым значением статуса возвращает 400."""
        r = self._patch_status(self.superuser, 'unknown_status')
        self.assertEqual(r.status_code, 400)

    def test_no_permission_returns_403(self):
        """Пользователь без view_task не проходит test_func и получает 403."""
        no_perm = make_user('no_perm_u', org=self.org)
        r = self._patch_status(no_perm, 'in_progress')
        self.assertEqual(r.status_code, 403)


# ---------------------------------------------------------------------------
# BugDetailView — POST: file upload
# ---------------------------------------------------------------------------

class BugDetailPostFileTest(ViewTestBase):
    def test_upload_file(self):
        """POST с файлом создаёт объект Attachment, привязанный к задаче."""
        self.client.force_login(self.task_user)
        f = SimpleUploadedFile('test.txt', b'file content', content_type='text/plain')
        r = self.client.post(
            self.url('signal1520:bug_details', pk=self.task.pk),
            {'file': f, 'description': 'Тест файл'},
        )
        self.assertEqual(r.status_code, 302)
        self.assertTrue(Attachment.objects.filter(task=self.task).exists())


# ---------------------------------------------------------------------------
# BugCreateView
# ---------------------------------------------------------------------------

class BugCreateViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на форму создания задачи перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:create_bug'))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_redirects_to_error(self):
        """Пользователь без add_task перенаправляется на страницу ошибки."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:create_bug'))
        self.assertEqual(r.status_code, 302)
        self.assertIn('error', r['Location'])

    def test_form_renders_for_permitted_user(self):
        """Пользователь с add_task видит форму создания задачи (GET 200)."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:create_bug'))
        self.assertEqual(r.status_code, 200)

    def test_form_contains_only_org_stations(self):
        """Выпадающий список station содержит только станции своей организации."""
        other_org = Organization.objects.create(name='Other', slug='other3')
        other_road = Road.objects.create(title='Дорога 2', organization=other_org)
        other_station = Station.objects.create(
            name='Чужая станция', road=other_road,
            created_by=self.superuser, organization=other_org,
        )
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:create_bug'))
        self.assertContains(r, 'Тест Станция')
        self.assertNotContains(r, 'Чужая станция')

    def test_create_bug_post(self):
        """POST с валидными данными создаёт новую задачу и редиректит на её страницу."""
        self.client.force_login(self.task_user)
        r = self.client.post(self.url('signal1520:create_bug'), {
            'station': self.station.pk,
            'description': 'Новая задача из теста',
            'responsible_organization': 'Тест Орг',
            'due_date': '',
        })
        self.assertEqual(r.status_code, 302)
        self.assertTrue(Task.objects.filter(description='Новая задача из теста').exists())

    def test_create_bug_sets_responsible_user(self):
        """form_valid() автоматически записывает в responsible_user текущего пользователя."""
        self.client.force_login(self.task_user)
        self.client.post(self.url('signal1520:create_bug'), {
            'station': self.station.pk,
            'description': 'Задача для проверки автора',
            'responsible_organization': '',
            'due_date': '',
        })
        task = Task.objects.get(description='Задача для проверки автора')
        self.assertEqual(task.responsible_user, self.task_user)

    def test_create_bug_missing_required_field(self):
        """POST без обязательных полей возвращает форму с ошибками, объект не создаётся."""
        self.client.force_login(self.task_user)
        r = self.client.post(self.url('signal1520:create_bug'), {'responsible_organization': 'Орг'})
        self.assertEqual(r.status_code, 200)


# ---------------------------------------------------------------------------
# BugUpdateView
# ---------------------------------------------------------------------------

class BugUpdateViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на форму редактирования задачи перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:bug_update', pk=self.task.pk))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_no_responsibility_returns_403(self):
        """Пользователь без change_task и не являющийся responsible_user получает 403."""
        other = make_user('other2', org=self.org)
        self.client.force_login(other)
        r = self.client.get(self.url('signal1520:bug_update', pk=self.task.pk))
        self.assertEqual(r.status_code, 403)

    def test_responsible_user_can_access(self):
        """Ответственный исполнитель может открыть форму редактирования своей задачи."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:bug_update', pk=self.task.pk))
        self.assertEqual(r.status_code, 200)

    def test_permitted_user_can_update(self):
        """POST с новыми данными обновляет задачу в БД для пользователя с change_task."""
        self.client.force_login(self.task_user)
        r = self.client.post(self.url('signal1520:bug_update', pk=self.task.pk), {
            'station': self.station.pk,
            'description': 'Обновлённая задача',
            'status': Task.Status.IN_PROGRESS,
            'responsible_organization': '',
        })
        self.assertEqual(r.status_code, 302)
        self.task.refresh_from_db()
        self.assertEqual(self.task.description, 'Обновлённая задача')

    def test_superuser_can_update(self):
        """Суперпользователь проходит test_func и получает форму редактирования."""
        self.client.force_login(self.superuser)
        r = self.client.get(self.url('signal1520:bug_update', pk=self.task.pk))
        self.assertEqual(r.status_code, 200)


# ---------------------------------------------------------------------------
# AlarmListView
# ---------------------------------------------------------------------------

class AlarmListViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на список алармов перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:alarm_list'))
        self.assertEqual(r.status_code, 302)

    def test_logged_in_returns_200(self):
        """Любой авторизованный пользователь видит справочник алармов."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:alarm_list'))
        self.assertEqual(r.status_code, 200)

    def test_alarm_appears_in_list(self):
        """Существующий аларм отображается в таблице на странице списка."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:alarm_list'))
        self.assertContains(r, 'T001')

    def test_search_by_number(self):
        """Поиск по номеру аларма возвращает соответствующую запись."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:alarm_list'), {'search': 'T001'})
        self.assertContains(r, 'T001')

    def test_search_no_result(self):
        """Поиск по несуществующему номеру не возвращает записей."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:alarm_list'), {'search': 'Z999'})
        self.assertNotContains(r, 'T001')


# ---------------------------------------------------------------------------
# KnowledgeCreateView
# ---------------------------------------------------------------------------

class KnowledgeCreateViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET перенаправляется на логин."""
        r = self.client.get(self.url('signal1520:knowledge_create'))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_redirects_to_error(self):
        """Пользователь без add_userknowledge перенаправляется на страницу ошибки."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:knowledge_create'))
        self.assertEqual(r.status_code, 302)
        self.assertIn('error', r['Location'])

    def test_with_permission_returns_200(self):
        """Пользователь с add_userknowledge видит форму добавления инструкции."""
        user = make_user('know_user', org=self.org, codenames=['add_userknowledge'])
        self.client.force_login(user)
        r = self.client.get(self.url('signal1520:knowledge_create'))
        self.assertEqual(r.status_code, 200)

    def test_superuser_can_access(self):
        """Суперпользователь проходит test_func без проверки прав."""
        self.client.force_login(self.superuser)
        r = self.client.get(self.url('signal1520:knowledge_create'))
        self.assertEqual(r.status_code, 200)


# ---------------------------------------------------------------------------
# KnowledgeDeleteView
# ---------------------------------------------------------------------------

class KnowledgeDeleteViewTest(ViewTestBase):
    def setUp(self):
        self.know_user = make_user(
            'know_del', org=self.org,
            codenames=['add_userknowledge', 'delete_userknowledge'],
        )
        knowledge = Knowledge.objects.create(
            created_by=self.know_user,
            external_link='https://example.com',
        )
        self.uk = UserKnowledge.objects.create(
            user=self.know_user,
            knowledge=knowledge,
            title='Тестовая инструкция',
        )

    def test_no_permission_returns_403_json(self):
        """Пользователь без delete_userknowledge получает JSON 403."""
        self.client.force_login(self.plain_user)
        r = self.client.post(self.url('signal1520:knowledge_delete', pk=self.uk.pk))
        self.assertEqual(r.status_code, 403)
        self.assertEqual(r.json()['error'], 'Недостаточно прав')

    def test_owner_with_permission_can_delete(self):
        """Владелец с delete_userknowledge может удалить свою инструкцию."""
        self.client.force_login(self.know_user)
        r = self.client.post(self.url('signal1520:knowledge_delete', pk=self.uk.pk))
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['success'])
        self.assertFalse(UserKnowledge.objects.filter(pk=self.uk.pk).exists())


# ---------------------------------------------------------------------------
# StationsExportView
# ---------------------------------------------------------------------------

class StationsExportViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на экспорт станций перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:stations_export'))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_returns_403(self):
        """Авторизованный пользователь без view_station получает 403."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:stations_export'))
        self.assertEqual(r.status_code, 403)

    def test_returns_xlsx_for_permitted_user(self):
        """Пользователь с view_station получает файл stations.xlsx."""
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:stations_export'))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            r['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('stations.xlsx', r['Content-Disposition'])

    def test_xlsx_contains_data(self):
        """Скачанный xlsx содержит данные существующей станции."""
        import io
        import openpyxl
        self.client.force_login(self.station_user)
        r = self.client.get(self.url('signal1520:stations_export'))
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        values = [str(cell.value) for row in wb.active.iter_rows() for cell in row]
        self.assertIn('Тест Станция', values)


# ---------------------------------------------------------------------------
# TasksExportView
# ---------------------------------------------------------------------------

class TasksExportViewTest(ViewTestBase):
    def test_redirect_anon(self):
        """Анонимный GET на экспорт задач перенаправляет на логин."""
        r = self.client.get(self.url('signal1520:bugs_export'))
        self.assertEqual(r.status_code, 302)

    def test_no_permission_returns_403(self):
        """Авторизованный пользователь без view_task получает 403."""
        self.client.force_login(self.plain_user)
        r = self.client.get(self.url('signal1520:bugs_export'))
        self.assertEqual(r.status_code, 403)

    def test_returns_xlsx_for_permitted_user(self):
        """Пользователь с view_task получает файл tasks.xlsx."""
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bugs_export'))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            r['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('tasks.xlsx', r['Content-Disposition'])

    def test_xlsx_contains_task_data(self):
        """Скачанный xlsx содержит описание существующей задачи."""
        import io
        import openpyxl
        self.client.force_login(self.task_user)
        r = self.client.get(self.url('signal1520:bugs_export'))
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        values = [str(cell.value) for row in wb.active.iter_rows() for cell in row]
        self.assertIn('Тестовая задача', values)
